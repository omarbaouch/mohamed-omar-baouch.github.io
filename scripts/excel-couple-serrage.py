# Calculateur Excel de couple de serrage (FR et EN), mêmes formules que le calculateur de
# /blog/couple-serrage-vis-tableau/ (VDI 2230 simplifiée, limites d'élasticité ISO 898-1).
# Les cellules contiennent des formules Excel : le classeur se recalcule à l'ouverture.
#   pip install openpyxl formulas
#   python3 scripts/excel-couple-serrage.py      → public/telechargements/…xlsx (+ contrôle)
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
# M : pas P, diamètre d'appui sous tête dw, diamètre du trou de passage dh (mm) — repris de la page
G = {3: (0.5, 4.6, 3.4), 4: (0.7, 5.9, 4.5), 5: (0.8, 6.9, 5.5), 6: (1, 8.9, 6.6), 8: (1.25, 11.6, 9),
     10: (1.5, 14.6, 11), 12: (1.75, 16.6, 13.5), 14: (2, 19.6, 15.5), 16: (2, 22.5, 17.5), 20: (2.5, 28.2, 22),
     24: (3, 33.6, 26), 27: (3, 38, 30), 30: (3.5, 42.8, 33), 36: (4, 51.1, 39)}
MU = [0.08, 0.10, 0.12, 0.14, 0.16, 0.20]
CLASSES = ['8.8', '10.9', '12.9']

T = {
    'fr': dict(
        fichier='calcul-couple-de-serrage-vis.xlsx', calc='Calculateur', tab='Tableau µ 0,12', data='Données',
        titre='Calcul du couple de serrage d\'une vis (VDI 2230, ISO 898-1)',
        sous='Choisissez la vis, la classe et le frottement dans les cellules jaunes.',
        vis='Vis (filetage métrique ISO à pas gros)', classe='Classe de qualité', mu='Coefficient de frottement µ (tête et filet)',
        res='Résultats', couple='Couple de serrage', precharge='Précontrainte de montage', unite_m='N·m', unite_f='kN',
        inter='Valeurs intermédiaires', pas='Pas P (mm)', d2='Diamètre sur flancs d2 (mm)', d3='Diamètre à fond de filet d3 (mm)',
        As='Section résistante As (mm²)', rp='Limite d\'élasticité Rp0,2 min (MPa)', dw='Diamètre d\'appui sous tête dw (mm)', dh='Trou de passage dh (mm)',
        hyp=('Hypothèses : méthode VDI 2230 simplifiée, utilisation de 90 % de la limite d\'élasticité (ν = 0,9), même frottement sous tête et dans le filet, '
             'limites d\'élasticité minimales de l\'ISO 898-1 (640 MPa en 8.8 jusqu\'à M16, 660 MPa au-delà ; 940 MPa en 10.9 ; 1 100 MPa en 12.9). '
             'µ = 0,12 correspond à une visserie acier huilée. Pour un assemblage critique, utilisez le coefficient mesuré ou celui du fabricant.'),
        source='Source et explications : https://www.baouch.fr/blog/couple-serrage-vis-tableau/',
        tab_titre='Couples de serrage (N·m) et précontraintes (kN), µ = 0,12', col_vis='Vis', col_pas='Pas',
    ),
    'en': dict(
        fichier='bolt-torque-calculator.xlsx', calc='Calculator', tab='Chart µ 0.12', data='Data',
        titre='Bolt tightening torque calculator (VDI 2230, ISO 898-1)',
        sous='Pick the bolt, the property class and the friction in the yellow cells.',
        vis='Bolt (ISO metric coarse thread)', classe='Property class', mu='Friction coefficient µ (head and thread)',
        res='Results', couple='Tightening torque', precharge='Assembly preload', unite_m='N·m', unite_f='kN',
        inter='Intermediate values', pas='Pitch P (mm)', d2='Pitch diameter d2 (mm)', d3='Minor diameter d3 (mm)',
        As='Stress area As (mm²)', rp='Minimum yield strength Rp0.2 (MPa)', dw='Head bearing diameter dw (mm)', dh='Clearance hole dh (mm)',
        hyp=('Assumptions: simplified VDI 2230 method, 90% utilisation of the yield strength (ν = 0.9), same friction under the head and in the thread, '
             'minimum yield strengths from ISO 898-1 (640 MPa for 8.8 up to M16, 660 MPa above; 940 MPa for 10.9; 1,100 MPa for 12.9). '
             'µ = 0.12 corresponds to oiled steel fasteners. For a critical joint, use the measured coefficient or the manufacturer\'s value.'),
        source='Source and explanations: https://www.baouch.fr/en/blog/couple-serrage-vis-tableau/',
        tab_titre='Tightening torques (N·m) and preloads (kN), µ = 0.12', col_vis='Bolt', col_pas='Pitch',
    ),
}

NOIR, JAUNE, GRIS = '15161A', 'FFF4C2', 'F3F3F5'
fin = Side(style='thin', color='CFD0D4')


def formules(d, P, dw, dh, rp, mu):
    """Expressions Excel (chaînes) pour une vis donnée par références de cellules."""
    d2 = f'({d}-0.64952*{P})'
    d3 = f'({d}-1.22687*{P})'
    ds = f'(({d2}+{d3})/2)'
    As = f'(PI()/4*{ds}^2)'
    k = f'(1.5*{d2}/{ds}*({P}/(PI()*{d2})+1.155*{mu}))'
    F = f'({As}*0.9*{rp}/SQRT(1+3*{k}^2))'
    M = f'({F}*(0.16*{P}+0.58*{d2}*{mu}+({dw}+{dh})/4*{mu})/1000)'
    return dict(d2=d2, d3=d3, As=As, F=F, M=M)


def classeur(lang):
    t = T[lang]
    wb = Workbook()
    calc = wb.active
    calc.title = t['calc']
    data = wb.create_sheet(t['data'])
    tab = wb.create_sheet(t['tab'])

    # --- données
    data.append(['M', 'd (mm)', t['pas'], t['dw'], t['dh']])
    for i, (m, (P, dw, dh)) in enumerate(G.items(), start=2):
        data.append([f'M{m}', m, P, dw, dh])
    data.append([])
    data.append(['µ'] + MU)
    data.append([t['classe']] + CLASSES)
    for c in data[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=NOIR)
    for col, w in zip('ABCDE', (8, 9, 14, 30, 24)):
        data.column_dimensions[col].width = w
    n = len(G) + 1  # dernière ligne de données

    # --- calculateur
    calc.column_dimensions['A'].width = 46
    calc.column_dimensions['B'].width = 16
    calc.column_dimensions['C'].width = 8
    calc['A1'] = t['titre']; calc['A1'].font = Font(bold=True, size=14)
    calc['A2'] = t['sous']; calc['A2'].font = Font(italic=True, color='555555')
    calc['A4'], calc['B4'] = t['vis'], 'M10'
    calc['A5'], calc['B5'] = t['classe'], '8.8'
    calc['A6'], calc['B6'] = t['mu'], 0.12
    for r in (4, 5, 6):
        calc[f'B{r}'].fill = PatternFill('solid', fgColor=JAUNE)
        calc[f'B{r}'].font = Font(bold=True)
        calc[f'B{r}'].border = Border(top=fin, bottom=fin, left=fin, right=fin)
        calc[f'B{r}'].alignment = Alignment(horizontal='center')
    calc['B6'].number_format = '0.00'
    dv = DataValidation(type='list', formula1=f"='{t['data']}'!$A$2:$A${n}", allow_blank=False)
    dv2 = DataValidation(type='list', formula1='"8.8,10.9,12.9"', allow_blank=False)
    dv3 = DataValidation(type='decimal', operator='between', formula1='0.04', formula2='0.3', allow_blank=False)
    dv3.error = 'µ : 0,04 – 0,30' if lang == 'fr' else 'µ: 0.04 – 0.30'
    for v, cell in ((dv, 'B4'), (dv2, 'B5'), (dv3, 'B6')):
        calc.add_data_validation(v); v.add(cell)

    look = lambda col: f"VLOOKUP($B$4,'{t['data']}'!$A$2:$E${n},{col},FALSE)"
    calc['A14'], calc['B14'] = 'd (mm)', f'={look(2)}'
    calc['A15'], calc['B15'] = t['pas'], f'={look(3)}'
    calc['A16'], calc['B16'] = t['dw'], f'={look(4)}'
    calc['A17'], calc['B17'] = t['dh'], f'={look(5)}'
    calc['A18'], calc['B18'] = t['rp'], '=IF($B$5="8.8",IF(B14<=16,640,660),IF($B$5="10.9",940,1100))'
    f = formules('B14', 'B15', 'B16', 'B17', 'B18', '$B$6')
    calc['A19'], calc['B19'] = t['d2'], '=' + f['d2']
    calc['A20'], calc['B20'] = t['d3'], '=' + f['d3']
    calc['A21'], calc['B21'] = t['As'], '=' + f['As']
    calc['A8'] = t['res']; calc['A8'].font = Font(bold=True, size=12)
    calc['A9'], calc['B9'], calc['C9'] = t['couple'], '=' + f['M'], t['unite_m']
    calc['A10'], calc['B10'], calc['C10'] = t['precharge'], '=' + f['F'] + '/1000', t['unite_f']
    for r in (9, 10):
        calc[f'B{r}'].font = Font(bold=True, size=13)
        calc[f'B{r}'].number_format = '0.0'
        calc[f'A{r}'].font = Font(bold=True)
    calc['A13'] = t['inter']; calc['A13'].font = Font(bold=True, color='555555')
    for r in range(14, 22):
        calc[f'B{r}'].number_format = '0.00' if r != 21 else '0.0'
        calc[f'A{r}'].font = Font(color='555555')
    calc['A23'] = t['hyp']; calc['A23'].alignment = Alignment(wrap_text=True, vertical='top')
    calc.merge_cells('A23:C23'); calc.row_dimensions[23].height = 105
    calc['A24'] = t['source']; calc['A24'].font = Font(color='B4501D')

    # --- tableau µ = 0,12 (formules, pour contrôle et impression)
    tab['A1'] = t['tab_titre']; tab['A1'].font = Font(bold=True, size=13)
    tete = [t['col_vis'], t['col_pas']] + [f'{c} N·m' for c in CLASSES] + [f'{c} kN' for c in CLASSES]
    tab.append([]); tab.append(tete)
    for c in tab[3]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=NOIR)
    for i, m in enumerate(G, start=4):
        r = i - 2  # ligne correspondante de « Données »
        dref = lambda col: f"'{t['data']}'!{col}{r}"
        tab[f'A{i}'] = f"={dref('A')}"
        tab[f'B{i}'] = f"={dref('C')}"
        for j, cl in enumerate(CLASSES):
            rp = ('IF(%s<=16,640,660)' % dref('B')) if cl == '8.8' else ('940' if cl == '10.9' else '1100')
            ff = formules(dref('B'), dref('C'), dref('D'), dref('E'), f'({rp})', '0.12')
            tab.cell(row=i, column=3 + j, value='=' + ff['M']).number_format = '0.0' if m <= 5 else '0'
            tab.cell(row=i, column=6 + j, value='=' + ff['F'] + '/1000').number_format = '0.0' if m <= 5 else '0'
        if i % 2:
            for col in range(1, 9):
                tab.cell(row=i, column=col).fill = PatternFill('solid', fgColor=GRIS)
    for col, w in zip('ABCDEFGH', (8, 7, 11, 11, 11, 10, 10, 10)):
        tab.column_dimensions[col].width = w
    tab.page_setup.orientation = 'landscape'
    wb.calculation.fullCalcOnLoad = True
    out = ROOT / 'public/telechargements' / ('' if lang == 'fr' else 'en') / t['fichier']
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def controle(path, lang):
    """Recalcule le classeur hors Excel et compare au tableau publié (M10 8.8 µ 0,12 = 48 N·m…)."""
    import formulas
    t = T[lang]
    xl = formulas.ExcelModel().loads(str(path)).finish()
    sol = xl.calculate()
    val = lambda sheet, cell: float(sol[f"'[{path.name}]{sheet.upper()}'!{cell}"].value[0][0])
    attendu = {(4, 'C'): 1.3, (9, 'C'): 48, (9, 'D'): 71, (9, 'E'): 83, (9, 'F'): 30, (17, 'E'): 4139, (17, 'H'): 730, (11, 'C'): 133, (12, 'C'): 206}
    for (row, col), v in attendu.items():
        x = val(t['tab'], f'{col}{row}')
        affiche = round(x, 1) if x < 10 else round(x)  # même arrondi que le tableau de la page
        assert affiche == v, f'{lang} {col}{row} : {x:.2f} ≠ {v}'
    m, f = val(t['calc'], 'B9'), val(t['calc'], 'B10')
    assert abs(m - 48) < 0.6 and abs(f - 30) < 0.6, (m, f)
    print(f'{lang} : contrôle OK (M10 8.8 µ 0,12 → {m:.1f} N·m, {f:.1f} kN)')


if __name__ == '__main__':
    for lang in ('fr', 'en'):
        p = classeur(lang)
        print(p.relative_to(ROOT), f'{p.stat().st_size // 1024} Ko')
        if '--sans-controle' not in sys.argv:
            controle(p, lang)
